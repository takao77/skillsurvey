"""Build the skill mapping workbook (Quant + Qual tab per job family) from submitted surveys.

    python build_workbook.py                           # out/Skill_Mapping_<wave>.xlsx
    python build_workbook.py --na-threshold 1.5        # Qual shows N/A when average importance < 1.5
    python build_workbook.py --include-raw             # adds individual answers (HR analytics only)

Quant tab (per level): N Raters, IMP (avg importance 0-4), PROF (avg proficiency 1-5, N/A excluded),
                       RUE (share answering "required upon entry", 0-1)
Qual tab  (per level): proficiency label = ROUND(PROF) -> Basic/Intermediate/Proficient/Advanced/Expert,
                       or N/A when IMP is below the threshold or most raters chose N/A.
"""
from __future__ import annotations

import argparse
import json
import re
import statistics as st
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from survey_core import BASE_DIR, WAVE, engine, load_library, managers, responses

LABELS = ["N/A", "Basic", "Intermediate", "Proficient", "Advanced", "Expert"]
FILLS = ["FFFFFF", "D9DEE8", "B4C7E7", "C6E0B4", "FFD966", "F4B183"]  # Wave 1 style colours
HEAD = PatternFill("solid", fgColor="1F2A44")
SUB = PatternFill("solid", fgColor="E7EBF3")
THIN = Side(style="thin", color="D0D5DF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def sheet_name(family: str, suffix: str) -> str:
    clean = re.sub(r"[\[\]:*?/\\]", "", family)
    return f"{clean[:31 - len(suffix) - 1]}-{suffix}"


def load_submissions():
    q = (select(managers.c.name, managers.c.email, managers.c.family, responses.c.submitted_data,
                responses.c.submitted_at)
         .select_from(managers.join(responses, managers.c.token == responses.c.token))
         .where((managers.c.wave == WAVE) & responses.c.submitted_data.is_not(None)))
    with engine.connect() as c:
        return [dict(r, data=json.loads(r["submitted_data"])) for r in c.execute(q).mappings()]


def aggregate(family: str, subs: list[dict]) -> dict:
    """{(skill, level): stats}"""
    fam = load_library()[family]
    mine = [s["data"] for s in subs if s["family"] == family]
    out = {}
    for sk in fam["skills"]:
        for lv in fam["levels"]:
            cells = [d.get(sk["name"], {}).get(lv["title"]) for d in mine]
            cells = [c for c in cells if c and c.get("imp") is not None]
            n = len(cells)
            profs = [c["prof"] for c in cells if c.get("prof")]  # 0 = N/A, excluded
            out[(sk["name"], lv["title"])] = {
                "n": n,
                "imp": st.mean(c["imp"] for c in cells) if n else None,
                "prof": st.mean(profs) if profs else None,
                "rue": st.mean(c["rue"] for c in cells) if n else None,
                "sd": st.stdev(profs) if len(profs) > 1 else 0.0,
                "na_share": (n - len(profs)) / n if n else None,
            }
    return out


def qual_label(a: dict, threshold: float) -> int:
    if not a["n"] or a["prof"] is None or a["imp"] < threshold or a["na_share"] > 0.5:
        return 0
    return min(5, max(1, int(a["prof"] + 0.5)))  # round half up, like Excel ROUND


def put_row(ws, values) -> int:
    """Write a row below the last used row (safe with merged header cells)."""
    r = ws.max_row + 1
    for i, v in enumerate(values, start=1):
        ws.cell(r, i, v)
    return r


def header_block(ws, family, levels, per_level, subheads):
    # row 1 = level titles, row 2 = sub headers
    for col, title in enumerate(["Role", "Skill Name", "Skill Definition"], start=1):
        ws.cell(1, col, title)
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
    for i, lv in enumerate(levels):
        c0 = 4 + i * per_level
        ws.cell(1, c0, f"{lv['title']} - {lv['grade']}")
        if per_level > 1:
            ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + per_level - 1)
        for j, sh in enumerate(subheads):
            ws.cell(2, c0 + j, sh)
        if per_level == 1:
            ws.merge_cells(start_row=1, start_column=c0, end_row=2, end_column=c0)
    last = 3 + len(levels) * per_level
    for r in (1, 2):
        for c in range(1, last + 1):
            cell = ws.cell(r, c)
            cell.font = Font(bold=True, color="FFFFFF" if r == 1 else "1F2A44")
            cell.fill = HEAD if r == 1 else SUB
            cell.alignment = CENTER
            cell.border = BORDER
    ws.row_dimensions[1].height = 34
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 52
    ws.freeze_panes = "D3"
    return last


def build(threshold: float, include_raw: bool, out: Path) -> Path:
    lib = load_library()
    subs = load_submissions()
    wb = Workbook()
    wb.remove(wb.active)
    flags = []

    with engine.connect() as c:
        invited = c.execute(select(managers.c.name, managers.c.email, managers.c.family, managers.c.token)
                            .where(managers.c.wave == WAVE)).mappings().all()
    submitted_emails = {(s["email"], s["family"]) for s in subs}

    for family, fam in lib.items():
        A = aggregate(family, subs)
        # ---- Quant
        ws = wb.create_sheet(sheet_name(family, "Quant"))
        last = header_block(ws, family, fam["levels"], 4, ["N Raters", "IMP", "PROF", "RUE"])
        for sk in fam["skills"]:
            row = [family, sk["name"], sk["definition"]]
            for lv in fam["levels"]:
                a = A[(sk["name"], lv["title"])]
                row += [a["n"], a["imp"], a["prof"], a["rue"]]
            put_row(ws, row)
        for r in ws.iter_rows(min_row=3, max_col=last):
            for cell in r:
                cell.border = BORDER
                cell.alignment = WRAP if cell.column <= 3 else CENTER
                if cell.column > 3 and (cell.column - 4) % 4 != 0:
                    cell.number_format = "0.00"
        for col in range(4, last + 1):
            ws.column_dimensions[get_column_letter(col)].width = 9

        # ---- Qual
        wq = wb.create_sheet(sheet_name(family, "Qual"))
        lastq = header_block(wq, family, fam["levels"], 1, [])
        for sk in fam["skills"]:
            labels = [qual_label(A[(sk["name"], lv["title"])], threshold) for lv in fam["levels"]]
            r = put_row(wq, [family, sk["name"], sk["definition"]] + [LABELS[i] for i in labels])
            for j, i in enumerate(labels):
                cell = wq.cell(r, 4 + j)
                cell.fill = PatternFill("solid", fgColor=FILLS[i])
                if i == 0:
                    cell.font = Font(color="8A93A6")
        for r in wq.iter_rows(min_row=3, max_col=lastq):
            for cell in r:
                cell.border = BORDER
                cell.alignment = WRAP if cell.column <= 3 else CENTER
        for col in range(4, lastq + 1):
            wq.column_dimensions[get_column_letter(col)].width = 18

        # ---- flags for HR review
        for sk in fam["skills"]:
            row = [A[(sk["name"], lv["title"])] for lv in fam["levels"]]
            for i in range(1, len(row)):
                a, b = row[i - 1], row[i]
                if a["prof"] is not None and b["prof"] is not None and b["prof"] < a["prof"] - 0.25:
                    flags.append([family, sk["name"], "Progression break",
                                  f"PROF drops from {fam['levels'][i-1]['title']} ({a['prof']:.2f}) "
                                  f"to {fam['levels'][i]['title']} ({b['prof']:.2f})"])
            for lv, a in zip(fam["levels"], row):
                if a["sd"] >= 1:
                    flags.append([family, sk["name"], "Low rater agreement",
                                  f"{lv['title']}: proficiency spread ±{a['sd']:.2f} across {a['n']} raters"])
            if row and row[0]["rue"] is not None and row[0]["rue"] >= 2 / 3:
                flags.append([family, sk["name"], "Required upon entry",
                              f"{row[0]['rue']:.0%} say it's needed on day one for {fam['levels'][0]['title']}"])

    # ---- Insights
    wi = wb.create_sheet("HR Insights", 0)
    wi.append(["Job Family", "Skill", "Flag", "Detail"])
    for f in flags:
        wi.append(f)
    # ---- Response status
    wr = wb.create_sheet("Response Status", 1)
    wr.append(["Manager", "Email", "Job Family", "Status"])
    for m in sorted(invited, key=lambda m: (m["family"], m["name"])):
        wr.append([m["name"], m["email"], m["family"],
                   "Submitted" if (m["email"], m["family"]) in submitted_emails else "Pending"])
    wr.append([])
    wr.append(["Response rate", f"{len(subs)}/{len(invited)}",
               f"{len(subs) / len(invited):.0%}" if invited else "-"])
    # ---- Raw (optional - individual answers, keep within HR analytics)
    if include_raw:
        wraw = wb.create_sheet("Raw Responses")
        wraw.append(["Manager", "Email", "Job Family", "Skill", "Job Level", "Importance",
                     "Proficiency (0 = N/A)", "Required Upon Entry", "Submitted (UTC)"])
        for s in subs:
            fam = lib[s["family"]]
            for sk in fam["skills"]:
                for lv in fam["levels"]:
                    c = s["data"].get(sk["name"], {}).get(lv["title"], {})
                    wraw.append([s["name"], s["email"], s["family"], sk["name"], lv["title"],
                                 c.get("imp"), c.get("prof"), c.get("rue"), s["submitted_at"]])
    for ws in (wi, wr) + ((wraw,) if include_raw else ()):
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEAD
        ws.freeze_panes = "A2"
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 60 if (ws is wi and col == 4) else 24

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"{len(subs)} submission(s) · {len(flags)} flag(s) · saved {out}")
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--na-threshold", type=float, default=1.5)
    ap.add_argument("--include-raw", action="store_true")
    ap.add_argument("--out", default=str(BASE_DIR / "out" / f"Skill_Mapping_{WAVE.replace(' ', '')}.xlsx"))
    a = ap.parse_args()
    build(a.na_threshold, a.include_raw, Path(a.out))
